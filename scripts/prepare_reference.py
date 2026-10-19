"""Build a clean, provenance-tagged reference directory from the supplied data."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "data" / "files"
OUTPUT_DIR = ROOT / "data" / "reference"
BOOKLET_PATH = ROOT / "data" / "health-facilities_booklet-2023-2024.xlsx"
GROUPS = ("O+", "O-", "A+", "A-", "B+", "B-", "AB+", "AB-")
SRC_NAME = "RBC Blood Transfusion Division Annual Report 2024"


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames is None:
            raise ValueError(f"CSV has no header: {path}")
        return list(reader.fieldnames), list(reader)


def write_csv(path: Path, fields: list[str], rows: Iterable[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def copy_tagged(
    filename: str, data_source: str, source_dir: Path, output_dir: Path
) -> list[dict[str, str]]:
    fields, rows = read_csv(source_dir / filename)
    if "data_source" in fields:
        raise ValueError(f"Unexpected pre-tagged input; review before replacing: {filename}")
    fields.append("data_source")
    for row in rows:
        row["data_source"] = data_source
    write_csv(output_dir / filename, fields, rows)
    return rows


def prepare_reference(source_dir: Path, output_dir: Path, booklet_path: Path) -> dict[str, int]:
    output_dir.mkdir(parents=True, exist_ok=True)
    row_counts = {
        "dim_district.csv": len(copy_tagged("dim_district.csv", "derived", source_dir, output_dir)),
        "dim_facility_osm.csv": len(copy_tagged("dim_facility_osm.csv", "real", source_dir, output_dir)),
        "district_pairs.csv": len(copy_tagged("district_pairs.csv", "derived", source_dir, output_dir)),
        "fact_blood_usage.csv": len(copy_tagged("fact_blood_usage.csv", "real", source_dir, output_dir)),
        "fact_hospitalisation.csv": len(
            copy_tagged("fact_hospitalisation.csv", "real", source_dir, output_dir)
        ),
    }

    district_rows = read_csv(output_dir / "dim_district.csv")[1]
    district_codes = {row["district"]: row["adm2_pcode"] for row in district_rows}
    workbook = load_workbook(booklet_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    referral_nodes = []
    for cells in sheet.iter_rows(min_row=2, values_only=True):
        province, district, sector, facility_name, level, ownership, facility_type = cells[:7]
        if province != "Kigali City" or facility_type != "Referral Hospital":
            continue
        normalized_district = str(district).strip().title()
        if normalized_district not in district_codes:
            raise ValueError(f"Unknown district in Ministry booklet: {normalized_district}")
        facility_name = str(facility_name).strip()
        slug = re.sub(r"[^a-z0-9]+", "-", facility_name.lower()).strip("-")
        referral_nodes.append(
            {
                "node_id": f"kigali-referral:{district_codes[normalized_district]}:{slug}",
                "province": province,
                "district": normalized_district,
                "district_code": district_codes[normalized_district],
                "sector": str(sector).strip(),
                "facility_name": facility_name,
                "facility_level": str(level).strip(),
                "facility_type": str(facility_type).strip(),
                "ownership": str(ownership).strip(),
                "source": "Rwanda Ministry of Health, Health Facilities Booklet 2023-2024",
                "data_source": "real",
            }
        )
    workbook.close()
    if not referral_nodes:
        raise ValueError("No Kigali referral hospitals were found in the supplied Ministry booklet")
    write_csv(
        output_dir / "kigali_referral_nodes.csv",
        [
            "node_id",
            "province",
            "district",
            "district_code",
            "sector",
            "facility_name",
            "facility_level",
            "facility_type",
            "ownership",
            "source",
            "data_source",
        ],
        referral_nodes,
    )
    row_counts["kigali_referral_nodes.csv"] = len(referral_nodes)

    matrix_fields, matrix_rows = read_csv(source_dir / "district_distance_km.csv")
    if matrix_fields[0] != "district":
        raise ValueError("Expected the first column of district_distance_km.csv to be the row index")
    matrix_districts = matrix_fields[1:]
    distances = []
    for row in matrix_rows:
        origin = row["district"]
        for destination in matrix_districts:
            distances.append(
                {
                    "from_district": origin,
                    "to_district": destination,
                    "straight_km": row[destination],
                    "data_source": "derived",
                }
            )
    write_csv(
        output_dir / "district_distance_km.csv",
        ["from_district", "to_district", "straight_km", "data_source"],
        distances,
    )
    row_counts["district_distance_km.csv"] = len(distances)

    collection_rows = [
        ("Kigali", "Kigali City", [13122, 1279, 6457, 411, 4953, 277, 1183, 77]),
        ("Butare", "South", [6743, 652, 3286, 214, 2595, 125, 496, 39]),
        ("Karongi", "West", [6580, 459, 3137, 156, 2463, 145, 530, 22]),
        ("Ruhengeri", "North", [6881, 542, 3748, 158, 2637, 143, 624, 18]),
        ("Rwamagana", "East", [6965, 499, 3379, 185, 2489, 141, 560, 13]),
    ]
    collections = []
    for rcbt, region, units_by_group in collection_rows:
        for group, units in zip(GROUPS, units_by_group):
            collections.append(
                {
                    "rcbt": rcbt,
                    "region": region,
                    "blood_group": group,
                    "units_collected": units,
                    "year": 2024,
                    "source": SRC_NAME + ", Table 1.9",
                    "data_source": "real",
                }
            )
    totals = {group: 0 for group in GROUPS}
    for row in collections:
        totals[str(row["blood_group"])] += int(row["units_collected"])
    if sum(totals.values()) != 84383:
        raise AssertionError("RBC collection transcription does not sum to 84,383")
    write_csv(
        output_dir / "rbc_2024_collected_by_group_rcbt.csv",
        ["rcbt", "region", "blood_group", "units_collected", "year", "source", "data_source"],
        collections,
    )
    row_counts["rbc_2024_collected_by_group_rcbt.csv"] = len(collections)

    indicators = [
        ("units_collected", 2024, 84383, "Table 1.6"),
        ("units_collected", 2023, 86812, "Table 1.6"),
        ("units_planned_collection", 2024, 89148, "Table 1.5"),
        ("hospital_demand_units_rcbt", 2024, 55852, "Table 3.1"),
        ("hospital_supply_units_rcbt", 2024, 55852, "Table 3.1"),
        ("hospital_demand_units_zipline", 2024, 71346, "Table 3.1"),
        ("hospital_supply_units_zipline", 2024, 70985, "Table 3.1"),
        ("hospital_demand_units_total", 2024, 127198, "Table 3.1"),
        ("hospital_supply_units_total", 2024, 126837, "Table 3.1"),
        ("unmet_demand_units_total", 2024, 361, "derived from Table 3.1"),
        ("blood_donors", 2024, 58688, "Table 1.8"),
    ]
    indicator_records = [
        {
            "indicator": name,
            "year": year,
            "value": value,
            "table": table,
            "source": SRC_NAME,
            "data_source": "derived" if name == "unmet_demand_units_total" else "real",
        }
        for name, year, value, table in indicators
    ]
    write_csv(
        output_dir / "rbc_2024_national_indicators.csv",
        ["indicator", "year", "value", "table", "source", "data_source"],
        indicator_records,
    )
    row_counts["rbc_2024_national_indicators.csv"] = len(indicator_records)

    indications = [
        ("Anemia", 19973, 1527, 276, 37),
        ("PPH (postpartum haemorrhage)", 8428, 2336, 2394, 1208),
        ("Bleeding", 3851, 813, 536, 222),
        ("Thrombocytopenia", 433, 3275, 123, 12),
        ("Sepsis", 1773, 992, 210, 46),
        ("Cancers", 1769, 918, 97, 30),
        ("Abortion", 1611, 124, 91, 29),
        ("Malaria", 875, 287, 16, 2),
        ("Uterine rupture", 344, 146, 147, 55),
    ]
    indication_records = []
    for indication, rbc, platelets, ffp, cryo in indications:
        indication_records.append(
            {
                "indication": indication,
                "rbc": rbc,
                "platelets": platelets,
                "ffp": ffp,
                "cryo": cryo,
                "total_components": rbc + platelets + ffp + cryo,
                "all_indication_total_components": 55006,
                "source": SRC_NAME + ", Table 4.3",
                "data_source": "real",
            }
        )
    write_csv(
        output_dir / "rbc_2024_transfusion_indications.csv",
        [
            "indication",
            "rbc",
            "platelets",
            "ffp",
            "cryo",
            "total_components",
            "all_indication_total_components",
            "source",
            "data_source",
        ],
        indication_records,
    )
    row_counts["rbc_2024_transfusion_indications.csv"] = len(indication_records)

    usage_rows = _read_usage_totals(output_dir)
    collected_total = sum(totals.values())
    usage_total = sum(usage_rows.values())
    comparison = []
    for group in GROUPS:
        usage_share = usage_rows[group] / usage_total
        collection_share = totals[group] / collected_total
        comparison.append(
            {
                "blood_group": group,
                "blood_bank_csv_units_used": usage_rows[group],
                "blood_bank_csv_share_pct": round(usage_share * 100, 1),
                "rbc_collected_2024_units": totals[group],
                "rbc_collected_share_pct": round(collection_share * 100, 1),
                "share_ratio_csv_vs_rbc": round(usage_share / collection_share, 2),
                "data_source": "derived",
                "use_restriction": "audit_only_not_for_model_or_simulation",
            }
        )
    write_csv(
        output_dir / "blood_group_mix_comparison.csv",
        [
            "blood_group",
            "blood_bank_csv_units_used",
            "blood_bank_csv_share_pct",
            "rbc_collected_2024_units",
            "rbc_collected_share_pct",
            "share_ratio_csv_vs_rbc",
            "data_source",
            "use_restriction",
        ],
        comparison,
    )
    row_counts["blood_group_mix_comparison.csv"] = len(comparison)
    return row_counts


def _read_usage_totals(output_dir: Path) -> dict[str, int]:
    usage_path = output_dir / "fact_blood_usage.csv"
    _, rows = read_csv(usage_path)
    totals = {group: 0 for group in GROUPS}
    for row in rows:
        group = row["blood_group"]
        totals[group] += int(row["units"])
    return totals


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=SOURCE_DIR)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--booklet", type=Path, default=BOOKLET_PATH)
    args = parser.parse_args()
    counts = prepare_reference(args.source_dir, args.output_dir, args.booklet)
    for filename, rows in sorted(counts.items()):
        print(f"{filename}: {rows} rows")
    print(f"Prepared {len(counts)} provenance-tagged reference CSV files in {args.output_dir}")


if __name__ == "__main__":
    main()
