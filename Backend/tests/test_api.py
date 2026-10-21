from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import main
from app.main import Settings, create_app
from app.storage import DatasetStore

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture()
def client(tmp_path: Path) -> TestClient:
    settings = Settings(
        api_key="test-api-key",
        upload_dir=tmp_path / "uploads",
        cors_origins=("http://localhost:5173",),
        rate_limit_per_minute=1000,
    )
    return TestClient(create_app(settings))


def _csv_upload(
    client: TestClient,
    content: bytes,
    filename: str = "readings.csv",
    headers: dict[str, str] | None = None,
) -> dict:
    response = client.post(
        "/api/v1/datasets",
        files={"file": (filename, content, "text/csv")},
        headers=headers or {"X-API-Key": "test-api-key"},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_health_and_auth_contract(client: TestClient):
    health = client.get("/health")
    assert health.status_code == 200
    assert health.json()["model_version"] == "v0.1.0"
    assert health.json()["data_source"] == "simulated"
    assert health.headers["x-request-id"]

    denied = client.get("/api/v1/reference/districts")
    assert denied.status_code == 401
    assert set(denied.json()) >= {"error", "data_source", "as_of"}
    assert denied.json()["error"]["code"] == "invalid_api_key"


def test_model_info_uses_frozen_bundle_and_metrics(client: TestClient):
    response = client.get(
        "/api/v1/model/info", headers={"X-API-Key": "test-api-key"}
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["model_version"] == "v0.1.0"
    assert body["trained_on"] == "simulated"
    assert body["metrics"]["data_source"] == "simulated"
    assert body["feature_list"]
    assert "re-validated on real data" in body["limitations"]


def test_reference_and_template_endpoints(client: TestClient):
    districts = client.get(
        "/api/v1/reference/districts", headers={"X-API-Key": "test-api-key"}
    )
    assert districts.status_code == 200
    assert len(districts.json()["districts"]) == 30
    assert districts.json()["districts"][0]["pcode"]

    template = client.get(
        "/api/v1/templates/readings.csv", headers={"X-API-Key": "test-api-key"}
    )
    assert template.status_code == 200
    assert template.text.startswith("timestamp,district,blood_group,stock_units")
    assert template.headers["x-data-source"] == "derived"

    schema = client.get(
        "/api/v1/templates/schema", headers={"X-API-Key": "test-api-key"}
    )
    assert schema.status_code == 200
    assert schema.json()["fields"][0]["name"] == "timestamp"

    inventory = client.get(
        "/api/v1/data-inventory", headers={"X-API-Key": "test-api-key"}
    )
    assert inventory.status_code == 200
    assert inventory.json()["rows"]


def test_openapi_exposes_typed_routes_and_api_key(client: TestClient):
    schema = client.get("/openapi.json").json()
    paths = schema["paths"]
    for path in (
        "/api/v1/model/info",
        "/api/v1/datasets",
        "/api/v1/datasets/{dataset_id}/mapping",
        "/api/v1/datasets/{dataset_id}/score",
        "/api/v1/datasets/{dataset_id}/results",
        "/api/v1/datasets/{dataset_id}/results/export.csv",
        "/api/v1/predict",
        "/api/v1/reference/districts",
        "/api/v1/readings",
        "/api/v1/inventory/summary",
        "/api/v1/stream",
    ):
        assert path in paths
    assert "APIKeyHeader" in schema["components"]["securitySchemes"]
    assert "DatasetUploadResponse" in schema["components"]["schemas"]
    assert "ErrorEnvelope" in schema["components"]["schemas"]


def test_upload_mapping_scoring_pagination_and_csv_export(client: TestClient, monkeypatch):
    content = (
        "timestamp,district,blood_group,stock_units,demand_units,supply_received\n"
        "2026-10-01,Huye,O+,12,2,3\n"
        "2026-10-02,Huye,O+,13,1,2\n"
        "2026-10-03,Huye,O+,14,1,2\n"
        "2026-10-04,Huye,O+,15,1,2\n"
        "2026-10-05,Huye,O+,16,1,2\n"
        "2026-10-06,Huye,O+,17,1,2\n"
        "2026-10-07,Huye,O+,18,1,2\n"
        "2026-10-08,Huye,O+,19,1,2\n"
        "2026-10-09,Huye,O+,20,1,2\n"
        "2026-10-10,Huye,O+,21,1,2\n"
        "2026-10-11,Huye,O+,22,1,2\n"
        "2026-10-12,Huye,O+,23,1,2\n"
        "2026-10-13,Huye,O+,24,1,2\n"
        "2026-10-14,Huye,O+,25,1,2\n"
    ).encode()
    uploaded = _csv_upload(client, content)
    assert uploaded["data_source"] == "entered"
    assert uploaded["validation_report"]["status"] == "suggested_mapping_preview"

    mapped = client.put(
        f"/api/v1/datasets/{uploaded['dataset_id']}/mapping",
        json={
            "mapping": {
                "timestamp": "timestamp",
                "district": "district",
                "blood_group": "blood_group",
                "stock_units": "stock_units",
                "demand_units": "demand_units",
                "supply_received": "supply_received",
            }
        },
        headers={"X-API-Key": "test-api-key"},
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["validation_report"]["accepted_rows"] == 14

    def fake_score(_frame, as_of=None, horizons=("24h", "48h"), top_k=5):
        return {
            "as_of": "2026-10-14T00:00:00Z",
            "model_version": "v0.1.0",
            "trained_on": "simulated",
            "data_source": "entered",
            "drift": {"distribution_shift": False, "psi_by_feature": {}},
            "results": [
                {
                    "district": "Huye",
                    "blood_group": "O+",
                    "risk_24h": 0.4,
                    "risk_48h": 0.5,
                    "rule_days_of_cover": 6.0,
                    "rank": 1,
                    "in_watchlist": True,
                    "top_drivers": ["low days of cover"],
                    "confidence": "medium",
                    "quality_flags": [],
                    "model_version": "v0.1.0",
                    "data_source": "entered",
                    "low_confidence_group": False,
                }
            ],
        }

    monkeypatch.setattr(main, "score", fake_score)
    scored = client.post(
        f"/api/v1/datasets/{uploaded['dataset_id']}/score",
        json={"top_k": 5, "horizons": ["24h", "48h"]},
        headers={"X-API-Key": "test-api-key"},
    )
    assert scored.status_code == 200, scored.text
    assert scored.json()["total"] == 1

    page = client.get(
        f"/api/v1/datasets/{uploaded['dataset_id']}/results?limit=1&confidence=medium",
        headers={"X-API-Key": "test-api-key"},
    )
    assert page.status_code == 200
    assert page.json()["results"][0]["top_drivers"] == ["low days of cover"]

    exported = client.get(
        f"/api/v1/datasets/{uploaded['dataset_id']}/results/export.csv",
        headers={"X-API-Key": "test-api-key"},
    )
    assert exported.status_code == 200
    assert "district,blood_group" in exported.text
    deleted = client.delete(
        f"/api/v1/datasets/{uploaded['dataset_id']}",
        headers={"X-API-Key": "test-api-key"},
    )
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] is True


def test_french_csv_duplicates_validation_rows_and_privacy_guard(client: TestClient):
    content = (
        "Rapport de stock\n"
        "jour;nom_district;groupe_sanguin;stock;consommation;livraison;patient_name\n"
        "2026-10-01;Huye;O POS;10;3;4;PRIVATE\n"
        "2026-10-02;Huye;Opos;11;3;4;PRIVATE\n"
        "2026-10-02;Huye;O Rh+;12;2;3;PRIVATE\n"
        "2026-10-03;Huye;A NEG;-1;1;0;PRIVATE\n"
        "2099-01-01;Huye;B+;5;1;0;PRIVATE\n"
    ).encode()
    uploaded = _csv_upload(client, content, "messy_french_semicolon.csv")
    assert uploaded["delimiter"] == ";"
    assert uploaded["preview"][0]["jour"] == "2026-10-01"
    assert "PRIVATE" not in str(uploaded)

    mapping = {
        key: uploaded["mapping_suggestions"][key]["column"]
        for key in ("timestamp", "district", "blood_group", "stock_units")
    }
    mapping.update(
        {
            "demand_units": "consommation",
            "supply_received": "livraison",
        }
    )
    normalized = client.put(
        f"/api/v1/datasets/{uploaded['dataset_id']}/mapping",
        json={"mapping": mapping},
        headers={"X-API-Key": "test-api-key"},
    )
    assert normalized.status_code == 200, normalized.text
    report = normalized.json()["validation_report"]
    assert report["duplicate_rows_removed"] == 1
    assert report["negative_unit_rows"] == 1
    assert report["future_date_rows"] == 1
    assert {"row_number": 6, "reasons": ["negative_stock_units"]} in report["rejected_rows"]
    assert {"row_number": 7, "reasons": ["future_timestamp"]} in report["rejected_rows"]


def test_excel_sheet_picker_and_document_preview_only(client: TestClient):
    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append(["Instructions"])
    stock = workbook.create_sheet("Stock")
    stock.append(["timestamp", "district", "blood_group", "stock_units"])
    stock.append(["2026-10-20", "Huye", "O+", 8])
    buffer = io.BytesIO()
    workbook.save(buffer)
    uploaded = _csv_upload(client, buffer.getvalue(), "messy_two_sheets.xlsx")
    assert uploaded["detected_sheet"] == "Stock"
    assert uploaded["available_sheets"] == ["README", "Stock"]

    docx = pytest.importorskip("docx")
    document = docx.Document()
    table = document.add_table(rows=2, cols=5)
    headers = ("timestamp", "district", "blood_group", "stock_units", "patient_name")
    values = ("2026-10-20", "Huye", "O+", "2", "PRIVATE")
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, values, strict=True):
        cell.text = value
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    preview = _csv_upload(client, docx_buffer.getvalue(), "table.docx")
    assert preview["preview_only"] is True
    assert "PRIVATE" not in str(preview)
    assert preview["extracted_tables"][0]["sensitive_columns_dropped"] == ["patient_name"]


def test_predict_json_validates_limit_and_retains_no_upload(client: TestClient, monkeypatch):
    def fake_score(_frame, as_of=None, horizons=("24h", "48h"), top_k=5):
        return {
            "as_of": "2026-10-20T00:00:00Z",
            "model_version": "v0.1.0",
            "trained_on": "simulated",
            "data_source": "entered",
            "drift": {"distribution_shift": False, "psi_by_feature": {}},
            "results": [],
        }

    monkeypatch.setattr(main, "score", fake_score)
    response = client.post(
        "/api/v1/predict",
        json={
            "rows": [
                {
                    "timestamp": "2026-10-20",
                    "district": "Huye",
                    "blood_group": "O+",
                    "stock_units": 8,
                }
            ]
        },
        headers={"X-API-Key": "test-api-key"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["data_source"] == "entered"
    assert list((client.app.state.store.root).iterdir()) == []


def test_short_history_no_demand_and_patient_column_api_fixtures(client: TestClient):
    short_history = "timestamp,district,blood_group,stock_units\n" + "".join(
        f"2026-10-{day:02d},Huye,O+,8\n" for day in range(1, 9)
    )
    short = _csv_upload(
        client, short_history.encode(), "short_history_8days.csv"
    )
    mapping = {
        "timestamp": "timestamp",
        "district": "district",
        "blood_group": "blood_group",
        "stock_units": "stock_units",
    }
    mapped_short = client.put(
        f"/api/v1/datasets/{short['dataset_id']}/mapping",
        json={"mapping": mapping},
        headers={"X-API-Key": "test-api-key"},
    )
    assert mapped_short.status_code == 200
    assert mapped_short.json()["validation_report"]["quality_flags_by_series"][
        "Huye|O+"
    ] == ["insufficient_history", "insufficient_inputs"]

    no_demand = "timestamp,district,blood_group,stock_units,supply_received\n" + "".join(
        f"2026-10-{day:02d},Huye,O+,{10 + day},2\n" for day in range(1, 15)
    )
    no_demand_upload = _csv_upload(
        client, no_demand.encode(), "no_demand_column.csv"
    )
    no_demand_mapping = dict(mapping)
    no_demand_mapping["supply_received"] = "supply_received"
    mapped_no_demand = client.put(
        f"/api/v1/datasets/{no_demand_upload['dataset_id']}/mapping",
        json={"mapping": no_demand_mapping},
        headers={"X-API-Key": "test-api-key"},
    )
    assert mapped_no_demand.status_code == 200
    assert mapped_no_demand.json()["validation_report"]["demand_estimated_rows"] > 0
    assert "demand_estimated" in mapped_no_demand.json()["validation_report"][
        "quality_flags_by_series"
    ]["Huye|O+"]

    privacy_content = (
        "timestamp,district,blood_group,stock_units,patient_name,national_id,phone,address\n"
        "2026-10-20,Huye,O+,4,PRIVATE,PRIVATE,PRIVATE,PRIVATE\n"
    ).encode()
    private_upload = _csv_upload(
        client, privacy_content, "with_patient_columns_FAKE.csv"
    )
    assert "PRIVATE" not in str(private_upload)
    assert set(private_upload["preview"][0]) == {
        "timestamp",
        "district",
        "blood_group",
        "stock_units",
    }
    assert "sensitive_columns_dropped" in str(private_upload["validation_report"])


def test_sample_30_day_upload_and_scoring_api(client: TestClient):
    sample_path = REPOSITORY_ROOT / "data" / "samples" / "sample_readings_30d_SIMULATED.csv"
    uploaded = _csv_upload(
        client, sample_path.read_bytes(), sample_path.name
    )
    mapping = {
        canonical: suggestion["column"]
        for canonical, suggestion in uploaded["mapping_suggestions"].items()
        if suggestion["column"]
    }
    confirmed = client.put(
        f"/api/v1/datasets/{uploaded['dataset_id']}/mapping",
        json={"mapping": mapping},
        headers={"X-API-Key": "test-api-key"},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["validation_report"]["series_count"] == 240

    scored = client.post(
        f"/api/v1/datasets/{uploaded['dataset_id']}/score",
        json={"top_k": 10, "horizons": ["24h", "48h"]},
        headers={"X-API-Key": "test-api-key"},
    )
    assert scored.status_code == 200, scored.text
    assert scored.json()["total"] == 240

    results = client.get(
        f"/api/v1/datasets/{uploaded['dataset_id']}/results?limit=240",
        headers={"X-API-Key": "test-api-key"},
    )
    assert results.status_code == 200
    assert len(results.json()["results"]) == 240
    assert all(
        row["confidence"] in {"medium", "high"}
        for row in results.json()["results"]
    )


def test_upload_size_limit_and_phase_four_stubs(client: TestClient, monkeypatch):
    monkeypatch.setattr(main, "MAX_UPLOAD_BYTES", 8)
    too_large = client.post(
        "/api/v1/datasets",
        files={"file": ("readings.csv", b"timestamp,district,blood_group,stock_units", "text/csv")},
        headers={"X-API-Key": "test-api-key"},
    )
    assert too_large.status_code == 413
    assert too_large.json()["error"]["code"] == "upload_too_large"
    for path, method in (
        ("/api/v1/readings", "post"),
        ("/api/v1/inventory/summary", "get"),
        ("/api/v1/stream", "get"),
    ):
        result = getattr(client, method)(path, headers={"X-API-Key": "test-api-key"})
        assert result.status_code == 501
        assert result.json()["error"]["code"] == "not_implemented"


def test_rate_limit_is_enforced(tmp_path: Path):
    app = create_app(
        Settings(
            api_key="test-api-key",
            upload_dir=tmp_path / "uploads",
            cors_origins=(),
            rate_limit_per_minute=1,
        )
    )
    with TestClient(app) as limited_client:
        assert limited_client.get("/health").status_code == 200
        limited = limited_client.get("/health")
    assert limited.status_code == 429
    assert limited.json()["error"]["code"] == "rate_limit_exceeded"


def test_retention_removes_only_expired_directories(tmp_path: Path):
    root = tmp_path / "uploads"
    store = DatasetStore(root)
    old_id = store.create(
        {
            "filename": "old.csv",
            "preview_only": False,
            "created_at": (datetime.now(UTC) - timedelta(days=8)).isoformat(),
        },
        pd.DataFrame({"safe": [1]}),
    )
    # The store assigns creation time; set the test fixture timestamp explicitly.
    metadata = store.get_metadata(old_id)
    metadata["created_at"] = (datetime.now(UTC) - timedelta(days=8)).isoformat()
    store.save_metadata(old_id, metadata)
    current_id = store.create(
        {"filename": "new.csv", "preview_only": False},
        pd.DataFrame({"safe": [2]}),
    )
    store.cleanup_expired()
    assert not (root / old_id).exists()
    assert (root / current_id).exists()


def test_application_startup_runs_retention_cleanup(tmp_path: Path):
    upload_dir = tmp_path / "uploads"
    store = DatasetStore(upload_dir)
    expired_id = store.create(
        {"filename": "expired.csv", "preview_only": False},
        pd.DataFrame({"safe": [1]}),
    )
    metadata = store.get_metadata(expired_id)
    metadata["created_at"] = (datetime.now(UTC) - timedelta(days=8)).isoformat()
    store.save_metadata(expired_id, metadata)
    app = create_app(
        Settings(
            api_key="test-api-key",
            upload_dir=upload_dir,
            cors_origins=(),
            rate_limit_per_minute=100,
        )
    )
    with TestClient(app) as startup_client:
        assert startup_client.get("/health").status_code == 200
    assert not (upload_dir / expired_id).exists()
